"""Real-data preparation, provenance, quality reports and leakage-aware splits."""

from collections import Counter, defaultdict
from dataclasses import asdict
from io import BytesIO
import csv
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import tempfile
from typing import Callable

import cv2
import numpy as np
from PIL import Image

from .archives import safe_destination, stage_source
from .artifacts import export_preview, file_sha256, result_json, runtime_versions, write_csv, write_json
from .audit import SPLIT_NAMES, SPLIT_RATIOS, assign_splits, audit_samples, leakage_check, perceptual_hashes
from .config import DEFAULT_CONFIG, FEATURE_SCHEMA, PREPROCESSING_VERSION, pipeline_fingerprint
from .features import FEATURE_NAMES
from .image_io import InputError, load_image, repair_png_profile_crc
from .normalization import extend_tissue, preprocess_with_details

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
MANIFEST_FIELDS = [
    "sample_id", "dataset", "cohort", "body_part", "image_role", "selected",
    "image_path", "original_image_path", "source_sha256", "pixel_sha256",
    "width", "height", "patient_id", "patient_id_verified", "source_group_id",
    "duplicate_group_id", "hb_g_dl", "hb_available", "anemia_label", "age_years", "sex",
    "is_valid", "inclusion_status", "exclusion_reason", "quality_warnings",
    "png_profile_crc_repair_required", "normalized_path", "normalized_sha256",
    "roi_mask_path", "roi_mask_sha256", "metadata_path", "metadata_sha256",
    "split", "split_basis", "preprocessing_version", "preprocessing_fingerprint",
]
FEATURE_FIELDS = ["sample_id", "normalized_path", "split", "split_basis", "hb_available",
                  "is_conjunctiva", "is_fingernail", "ita_angle", "ita_category", "glare_fraction",
                  "preprocessing_fingerprint", *FEATURE_NAMES]


def parse_number(value) -> float | None:
    if value is None or (isinstance(value, str) and value.strip().lower() in {"", "_", "n/a", "na", "nan"}):
        return None
    number = float(str(value).strip().replace(",", "."))
    if not math.isfinite(number):
        return None
    return number


def patient_number(value) -> str:
    number = parse_number(value)
    if number is None or not number.is_integer() or number <= 0:
        raise ValueError(f"Invalid patient identifier: {value!r}")
    return str(int(number))


def read_eye_labels(path: Path) -> dict[str, dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Spreadsheet import requires pip install 'vyra-preprocessing[data]'") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = iter(workbook.active.iter_rows(values_only=True))
        header = [str(v).strip().lower() if v is not None else "" for v in next(rows)]
        if not {"number", "hgb", "gender", "age"}.issubset(header):
            raise ValueError(f"Missing required spreadsheet columns in {path.name}")
        records = {}
        for values in rows:
            row = dict(zip(header, values))
            if row.get("number") is None:
                continue
            identifier = patient_number(row["number"])
            if identifier in records:
                raise ValueError(f"Duplicate patient metadata row {identifier} in {path.name}")
            hb = parse_number(row.get("hgb"))
            if hb is not None and hb <= 0:
                raise ValueError("Laboratory Hb must be positive or explicitly missing")
            records[identifier] = {
                "hb_g_dl": hb,
                "age_years": parse_number(row.get("age")),
                "sex": str(row.get("gender") or "").strip().upper(),
            }
        return records
    finally:
        workbook.close()


def nail_identity(filename: str) -> tuple[str, int]:
    stem = re.sub(r" \(\d+\)$", "", Path(filename).stem)
    match = re.fullmatch(r"([A-Za-z-]+)-(\d+)(?:FV|RA)?(?:\s+\d+)*", stem, flags=re.IGNORECASE)
    if not match:
        raise ValueError(f"Unrecognized nail filename: {filename}")
    prefix = match[1].lower()
    labels = {"anemic-fin": 1, "anemic-fn": 1, "anmeic-fn": 1,
              "non-anemic-fin": 0, "non-anrmic-fn": 0}
    if prefix not in labels:
        raise ValueError(f"Unrecognized nail label prefix: {prefix}")
    label = labels[prefix]
    # Conservative image-family identifier, NOT a claim of patient identity.
    return f"ghana:source:{label}:{int(match[2])}", label


def eye_role(filename: str) -> str:
    if Path(filename).suffix.lower() in {".jpg", ".jpeg"}:
        return "full_eye"
    stem = re.sub(r"\(\d+\)$", "", Path(filename).stem.lower())
    stem = stem.replace("papebral", "palpebral").replace("palplebral", "palpebral")
    if stem.endswith("_forniceal_palpebral"):
        return "forniceal_palpebral"
    if stem.endswith("_palpebral"):
        return "palpebral"
    if stem.endswith("_forniceal"):
        return "forniceal"
    return "unknown"


def segmented_eye_mask(rgb: np.ndarray) -> np.ndarray:
    """Only for known EDA segmented images: remove border-connected pure white."""
    white = np.all(rgb == 255, axis=2).astype(np.uint8)
    _, labels = cv2.connectedComponents(white, connectivity=4)
    border = np.unique(np.concatenate([labels[0], labels[-1], labels[:, 0], labels[:, -1]]))
    border = border[border != 0]
    background = np.isin(labels, border) & white.astype(bool)
    return ~background


def source_signature(source: Path) -> dict:
    source = source.resolve()
    if source.is_file():
        return {"name": source.name, "sha256": file_sha256(source), "kind": "archive"}
    if not source.is_dir():
        raise ValueError(f"Dataset source is missing: {source}")
    digest = hashlib.sha256()
    count = 0
    for path in sorted(source.rglob("*")):
        if path.is_symlink():
            raise ValueError("Dataset directory must not contain symlinks")
        if path.is_file():
            digest.update(path.relative_to(source).as_posix().encode())
            digest.update(file_sha256(path).encode())
            count += 1
    return {"name": source.name, "sha256": digest.hexdigest(), "kind": "directory", "files": count}


def _preparation_fingerprint(eyes: dict, nails: dict, seed: int, radius: int) -> str:
    code = hashlib.sha256()
    for name in ("datasets.py", "archives.py", "audit.py", "artifacts.py"):
        code.update(Path(__file__).with_name(name).read_bytes())
    payload = {"eyes": eyes, "nails": nails, "seed": seed, "phash_radius": radius,
               "preprocessing": pipeline_fingerprint(), "importer_sha256": code.hexdigest()}
    return hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()


def _asset(path: Path, root: Path, **attributes) -> dict:
    relative = path.relative_to(root).as_posix()
    record = dict.fromkeys(MANIFEST_FIELDS)
    record.update({"sample_id": "asset:" + hashlib.sha256(relative.encode()).hexdigest()[:24],
                   "image_path": relative, "selected": False, "patient_id_verified": False,
                   "hb_available": False, "inclusion_status": "not_selected", "quality_warnings": [],
                   "preprocessing_version": PREPROCESSING_VERSION,
                   "preprocessing_fingerprint": pipeline_fingerprint()})
    record.update(attributes)
    data = path.read_bytes()
    record["source_sha256"] = hashlib.sha256(data).hexdigest()
    try:
        repaired, repair_needed = repair_png_profile_crc(data)
        record["png_profile_crc_repair_required"] = repair_needed
        with Image.open(BytesIO(repaired)) as image:
            record["width"], record["height"] = image.size
    except (InputError, OSError, ValueError, SyntaxError) as exc:
        record["_decode_error"] = str(exc)
    return record


def inventory_eyes(directory: Path, root: Path) -> list[dict]:
    assets = []
    for country in ("India", "Italy"):
        workbooks = sorted(directory.rglob(country + ".xlsx"))
        if len(workbooks) != 1:
            raise ValueError(f"Expected exactly one {country}.xlsx spreadsheet, found {len(workbooks)}")
        labels = read_eye_labels(workbooks[0])
        cohort_root = workbooks[0].parent
        found = set()
        for folder in sorted((p for p in cohort_root.iterdir() if p.is_dir()), key=lambda p: p.name):
            identifier = patient_number(folder.name)
            if identifier not in labels:
                raise ValueError(f"Image folder {country}/{identifier} has no metadata row")
            found.add(identifier)
            patient_id = f"eda:{country.lower()}:{identifier}"
            family = []
            for path in sorted(folder.iterdir()):
                if path.suffix.lower() in IMAGE_EXTENSIONS and path.is_file():
                    row = _asset(path, root, dataset="eyes_defy_anemia", cohort=country.lower(),
                                 body_part="conjunctiva", image_role=eye_role(path.name),
                                 patient_id=patient_id, patient_id_verified=True, source_group_id=patient_id,
                                 anemia_label=None, **labels[identifier])
                    row["hb_available"] = row["hb_g_dl"] is not None
                    family.append(row)
            originals = [r for r in family if r["image_role"] == "full_eye"]
            candidates = [r for r in family if r["image_role"] == "palpebral"]
            if not originals or not candidates:
                raise ValueError(f"Patient {patient_id} lacks an original image or palpebral crop")
            selected = min(candidates, key=lambda r: (bool(re.search(r"\(\d+\)", r["image_path"])), r["image_path"]))
            selected["selected"] = True
            selected["original_image_path"] = originals[0]["image_path"]
            assets.extend(family)
        if found != set(labels):
            raise ValueError(f"{country} metadata contains patients without image folders")
    return assets


def inventory_nails(directory: Path, root: Path) -> list[dict]:
    assets = []
    for path in sorted(p for p in directory.rglob("*") if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS):
        group, label = nail_identity(path.name)
        assets.append(_asset(path, root, dataset="ghana_fingernails", cohort="ghana", body_part="fingernail",
                             image_role="nail", selected=True, patient_id=None, patient_id_verified=False,
                             source_group_id=group, anemia_label=label, hb_g_dl=None, hb_available=False))
    if not assets:
        raise ValueError("Nail source contains no recognized images")
    return assets


def _process_samples(assets: list[dict], root: Path, progress: Callable[[str], None]) -> list[dict]:
    samples = [row for row in assets if row["selected"]]
    examples = set()
    (root / "derived").mkdir()
    for index, row in enumerate(samples, 1):
        source = root / row["image_path"]
        try:
            loaded = load_image(source)
            pixel_bytes = np.ascontiguousarray(loaded.rgb).tobytes() + str(loaded.rgb.shape).encode()
            if loaded.alpha_mask is not None:
                pixel_bytes += np.packbits(loaded.alpha_mask).tobytes()
            row["pixel_sha256"] = hashlib.sha256(pixel_bytes).hexdigest()
            mask = loaded.alpha_mask
            if row["dataset"] == "eyes_defy_anemia" and mask is None:
                mask = segmented_eye_mask(loaded.rgb)
            if mask is None:
                mask = np.ones(loaded.rgb.shape[:2], dtype=bool)
            details = preprocess_with_details(loaded.rgb, row["body_part"], roi_mask=mask, pre_cropped=True)
            details.metadata.update(loaded.metadata)
            details.metadata["source_asset_id"] = row["sample_id"]
            details.metadata["original_image_path"] = row.get("original_image_path")
            row["is_valid"] = details.result.is_valid
            row["quality_warnings"] = details.result.quality_warnings.copy()
            if loaded.metadata["png_profile_crc_repaired"]:
                row["quality_warnings"].append("PNG_PROFILE_CRC_REPAIRED: source unchanged")
                details.result.quality_warnings = row["quality_warnings"].copy()
            row["width"], row["height"] = loaded.rgb.shape[1], loaded.rgb.shape[0]
            if not details.result.is_valid:
                row["inclusion_status"] = "invalid"
                row["exclusion_reason"] = details.metadata["error_code"]
            else:
                row["inclusion_status"] = "ready" if row["hb_available"] or row["anemia_label"] is not None else "unlabeled"
                row["exclusion_reason"] = "" if row["inclusion_status"] == "ready" else "missing_laboratory_hb"
                row["_features"] = details.result.feature_vector.tolist()
                row["_ita_angle"] = details.result.ita_angle
                row["_ita_category"] = details.result.ita_category
                row["_glare_fraction"] = details.result.glare_fraction
                row["_body_encoding"] = details.result.body_part_encoding.tolist()
                row["_phashes"] = perceptual_hashes(extend_tissue(details.raw_roi_rgb, details.roi_mask))
                base = row["sample_id"].replace(":", "-")
                normalized = root / "derived" / (base + ".png")
                mask_path = root / "derived" / (base + ".mask.png")
                sidecar = root / "derived" / (base + ".json")
                Image.fromarray(details.result.normalized_image_rgb).save(normalized)
                Image.fromarray(details.roi_mask.astype(np.uint8) * 255).save(mask_path)
                write_json(sidecar, result_json(details))
                for key, path in (("normalized", normalized), ("roi_mask", mask_path), ("metadata", sidecar)):
                    row[key + "_path"] = path.relative_to(root).as_posix()
                    row[key + "_sha256"] = file_sha256(path)
            example = row["body_part"] if details.result.is_valid else "invalid-" + details.metadata["error_code"]
            if example not in examples:
                export_preview(details, root / "examples" / example)
                examples.add(example)
        except InputError as exc:
            row["is_valid"] = False
            row["inclusion_status"] = "invalid"
            row["exclusion_reason"] = exc.code
            row["quality_warnings"] = [f"{exc.code}: {exc}"]
        if index % 200 == 0 or index == len(samples):
            progress(f"Preprocessed {index}/{len(samples)} selected samples")
    return samples


def _summary(assets: list[dict], samples: list[dict], events: list[dict], seed: int, radius: int) -> dict:
    included = [r for r in samples if r["split"]]
    partitions = {}
    for dataset in ("eyes_defy_anemia", "ghana_fingernails"):
        rows = [r for r in samples if r["dataset"] == dataset]
        partitions[dataset] = {
            "selected_before_quality_audit": len(rows),
            "status_counts": dict(Counter(r["inclusion_status"] for r in rows)),
            "exclusion_counts": dict(Counter(r["exclusion_reason"] for r in rows if r["exclusion_reason"])),
            "splits": {},
        }
        for split in SPLIT_NAMES:
            subset = [r for r in rows if r["split"] == split]
            partitions[dataset]["splits"][split] = {
                "images": len(subset),
                "source_groups": len({r["source_group_id"] for r in subset}),
                "duplicate_components": len({r["duplicate_group_id"] for r in subset}),
                "verified_patients": len({r["patient_id"] for r in subset if r["patient_id_verified"]}),
                "with_lab_hb": sum(r["hb_available"] for r in subset),
                "source_class_counts": dict(Counter(str(r["anemia_label"]) for r in subset if r["anemia_label"] is not None)),
                "cohort_counts": dict(Counter(r["cohort"] for r in subset)),
            }
    return {
        "preprocessing_version": PREPROCESSING_VERSION,
        "preprocessing_fingerprint": pipeline_fingerprint(),
        "feature_schema": FEATURE_SCHEMA, "feature_names": list(FEATURE_NAMES),
        "seed": seed, "requested_split_ratios": dict(zip(SPLIT_NAMES, SPLIT_RATIOS)),
        "perceptual_hash_hamming_radius": radius,
        "inventory_image_count": len(assets),
        "inventory_by_dataset": dict(Counter(r["dataset"] for r in assets)),
        "eye_patient_count": len({r["patient_id"] for r in assets if r["patient_id"]}),
        "eye_patients_with_lab_hb": len({r["patient_id"] for r in assets if r["patient_id"] and r["hb_available"]}),
        "png_profile_crc_repair_required_assets": sum(bool(r["png_profile_crc_repair_required"]) for r in assets),
        "included_image_count": len(included), "datasets": partitions,
        "duplicate_audit_event_counts": dict(Counter(e["kind"] for e in events)),
        "leakage": leakage_check(samples),
        "runtime": runtime_versions(),
        "limitations": [
            "Nail archive supplies binary class labels but no laboratory Hb or verified patient mapping.",
            "Nail splits are exploratory source-image-group partitions, not verified patient-level evaluation.",
            "Perceptual matches are candidates, not proof that two patients are the same person.",
            "Eye classification labels are not invented; Person 2 owns demographic/clinical label policy.",
            "ITA requires an explicit surrounding-skin mask; dataset crops normally report Unknown.",
            "The supplied Person 2 note's counts and prior partitions could not be verified without its CSVs.",
            "CP-AnemiC is not supplied and is not part of these manifests.",
            "No model accuracy, clinical performance or Hb prediction is measured by this preprocessing package.",
        ],
    }


def _write_outputs(root: Path, assets: list[dict], samples: list[dict], events: list[dict], seed: int, radius: int) -> dict:
    samples = sorted(samples, key=lambda r: r["sample_id"])
    assets = sorted(assets, key=lambda r: r["image_path"])
    write_csv(root / "manifests" / "dataset_inventory.csv", assets, MANIFEST_FIELDS)
    write_csv(root / "manifests" / "sample_manifest.csv", samples, MANIFEST_FIELDS)
    write_csv(root / "manifests" / "duplicate_audit.csv", events, ["kind", "left_sample_id", "right_sample_id", "distance"])
    for split in SPLIT_NAMES:
        rows = [r for r in samples if r["split"] == split]
        write_csv(root / "splits" / (split + ".csv"), rows, MANIFEST_FIELDS)
        write_csv(root / "verified_splits" / (split + ".csv"), [r for r in rows if r["patient_id_verified"]], MANIFEST_FIELDS)
    feature_rows = []
    for row in samples:
        if not row["split"]:
            continue
        feature_rows.append({**{key: row.get(key) for key in FEATURE_FIELDS},
                             "is_conjunctiva": row["_body_encoding"][0], "is_fingernail": row["_body_encoding"][1],
                             "ita_angle": row["_ita_angle"] if math.isfinite(row["_ita_angle"]) else None,
                             "ita_category": row["_ita_category"], "glare_fraction": row["_glare_fraction"],
                             **dict(zip(FEATURE_NAMES, row["_features"]))})
    write_csv(root / "manifests" / "features.csv", feature_rows, FEATURE_FIELDS)
    summary = _summary(assets, samples, events, seed, radius)
    write_json(root / "reports" / "audit_summary.json", summary)
    write_json(root / "reports" / "configuration.json", asdict(DEFAULT_CONFIG))
    return summary


def prepare_datasets(eyes_source: Path, nails_source: Path, output: Path, *, seed: int = 42,
                     phash_radius: int = 6, progress: Callable[[str], None] = print) -> dict:
    if not 0 <= phash_radius <= 12:
        raise ValueError("pHash radius must be between 0 and 12")
    output = output.resolve()
    for source in (eyes_source, nails_source):
        if source.resolve().is_relative_to(output) or output.is_relative_to(source.resolve()) and source.is_dir():
            raise ValueError("Output directory and source dataset directories must not contain one another")
    progress("Hashing source archives/directories")
    eyes_signature, nails_signature = source_signature(eyes_source), source_signature(nails_source)
    preparation_id = _preparation_fingerprint(eyes_signature, nails_signature, seed, phash_radius)
    if output.exists():
        run_file = output / "run.json"
        if run_file.is_file() and json.loads(run_file.read_text())["preparation_fingerprint"] == preparation_id:
            progress("Matching prepared dataset exists; verifying before reuse")
            verification = verify_prepared(output)
            if not verification["passed"]:
                raise ValueError("Existing prepared data failed verification; use a new output directory")
            return json.loads((output / "reports" / "audit_summary.json").read_text())
        raise ValueError("Output already exists with a different or incomplete run; select a new output directory")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".vyra-prepare-", dir=output.parent) as temporary:
        root = Path(temporary) / "prepared"
        root.mkdir()
        progress("Extracting eye ZIP and nail ZIP/RAR into staging")
        stage_source(eyes_source, root / "raw" / "eyes")
        stage_source(nails_source, root / "raw" / "nails", nails=True)
        progress("Reading spreadsheet labels and inventorying source images")
        assets = inventory_eyes(root / "raw" / "eyes", root) + inventory_nails(root / "raw" / "nails", root)
        samples = _process_samples(assets, root, progress)
        progress("Auditing byte/pixel duplicates and perceptual candidates")
        events = audit_samples(samples, phash_radius)
        assign_splits(samples, seed)
        checks = leakage_check(samples)
        if not checks["passed"]:
            raise RuntimeError("Internal split validation failed; no prepared run was published")
        if any(not any(r["split"] == split and r["patient_id_verified"] for r in samples) for split in SPLIT_NAMES):
            raise ValueError("Insufficient verified eye groups for nonempty train/validation/test partitions")
        summary = _write_outputs(root, assets, samples, events, seed, phash_radius)
        artifacts = {p.relative_to(root).as_posix(): file_sha256(p) for directory in ("manifests", "splits", "verified_splits", "reports") for p in sorted((root / directory).glob("*")) if p.is_file()}
        write_json(root / "run.json", {
            "preparation_fingerprint": preparation_id,
            "preprocessing_version": PREPROCESSING_VERSION,
            "preprocessing_fingerprint": pipeline_fingerprint(),
            "sources": {"eyes": eyes_signature, "nails": nails_signature},
            "seed": seed, "phash_radius": phash_radius, "runtime": runtime_versions(),
            "artifact_sha256": artifacts,
        })
        root.rename(output)
    progress(f"Prepared {summary['included_image_count']} included samples; artifacts published atomically")
    return summary


def _read_csv(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def verify_prepared(root: Path) -> dict:
    root = root.resolve()
    errors = []
    try:
        run = json.loads((root / "run.json").read_text())
        if run["preprocessing_fingerprint"] != pipeline_fingerprint():
            errors.append("Preprocessing code/config/runtime fingerprint differs from this prepared run")
        for relative, expected in run["artifact_sha256"].items():
            path = safe_destination(root, relative)
            if not path.is_file() or file_sha256(path) != expected:
                errors.append(f"Artifact hash mismatch: {relative}")
        rows = _read_csv(root / "manifests" / "dataset_inventory.csv")
        for row in rows:
            source = safe_destination(root, row["image_path"])
            if not source.is_file() or file_sha256(source) != row["source_sha256"]:
                errors.append(f"Source image mismatch: {row['sample_id']}")
            for name in ("normalized", "roi_mask", "metadata"):
                if row.get(name + "_path"):
                    path = safe_destination(root, row[name + "_path"])
                    if not path.is_file() or file_sha256(path) != row[name + "_sha256"]:
                        errors.append(f"Derived artifact mismatch: {row['sample_id']} {name}")
        samples = _read_csv(root / "manifests" / "sample_manifest.csv")
        checks = leakage_check(samples)
        if not checks["passed"]:
            errors.append("Split overlap detected")
        expected_ids = {r["sample_id"] for r in samples if r["split"]}
        all_partition_ids = []
        for split in SPLIT_NAMES:
            partition = _read_csv(root / "splits" / (split + ".csv"))
            if any(r["split"] != split for r in partition):
                errors.append(f"Incorrect partition assignment in {split}")
            all_partition_ids.extend(r["sample_id"] for r in partition)
            verified = _read_csv(root / "verified_splits" / (split + ".csv"))
            if any(r["patient_id_verified"] != "true" or r["split"] != split for r in verified):
                errors.append(f"Unverified identity in verified {split} partition")
        if set(all_partition_ids) != expected_ids or len(all_partition_ids) != len(expected_ids):
            errors.append("Partition sample IDs do not match the included sample manifest")
        features = _read_csv(root / "manifests" / "features.csv")
        if {r["sample_id"] for r in features} != expected_ids or len(features) != len(expected_ids):
            errors.append("Feature rows do not match included sample IDs")
        for row in features:
            if any(not math.isfinite(float(row[name])) for name in FEATURE_NAMES):
                errors.append(f"Nonfinite feature in {row['sample_id']}")
    except (OSError, ValueError, KeyError, csv.Error) as exc:
        errors.append(f"Verification could not complete: {exc}")
        checks = {}
        rows = []
    return {"passed": not errors, "errors": errors, "inventory_images_checked": len(rows), "leakage": checks}


def export_handoff(root: Path, destination: Path) -> None:
    root, destination = root.resolve(), destination.resolve()
    if destination.is_relative_to(root) or root.is_relative_to(destination):
        raise ValueError("Handoff output and prepared dataset directories must not contain one another")
    verification = verify_prepared(root)
    if not verification["passed"]:
        raise ValueError("Cannot export an unverified prepared dataset: " + "; ".join(verification["errors"]))
    destination.mkdir(parents=True, exist_ok=True)
    for name in ("manifests", "splits", "verified_splits", "reports"):
        shutil.copytree(root / name, destination / name, dirs_exist_ok=True)
    shutil.copyfile(root / "run.json", destination / "run.json")
    write_json(destination / "verification.json", verification)
    (destination / "README.md").write_text(
        "# Observed data handoff\n\nThese files contain metadata/features only; no photographs are included.\n"
        "All image and derived paths are relative to the dataset preparation output directory, not this report folder.\n"
        "Run preparation from the same source archives and pinned environment to reproduce the artifacts.\n"
        "`splits/` includes exploratory nail groups; `verified_splits/` contains verified eye patients only.\n"
        "Read `reports/audit_summary.json` for observed exclusions and limitations.\n",
        encoding="utf-8")
